#!/usr/bin/env python
"""PDF tables -> clean Excel workbook.

Usage:  python convert.py input.pdf output.xlsx
- Finds every table on every page (PyMuPDF table detection), one sheet per table, plus a "Summary" sheet.
- Normalises whitespace, converts numeric strings to numbers, strips currency symbols/thousand separators.
- Flags duplicate rows and empty cells on a "QA" sheet so the client sees what needs a decision.
"""
import re
import sys

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

NUM = re.compile(r"^[\s$€£₹]*[-+]?\d[\d,]*(\.\d+)?\s*%?$")


def clean(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    if NUM.match(s):
        n = re.sub(r"[^\d.\-+]", "", s)
        try:
            return float(n) if "." in n else int(n)
        except ValueError:
            return s
    return s


def convert(pdf_path, xlsx_path):
    doc = fitz.open(pdf_path)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Sheet", "Page", "Rows", "Columns"])
    qa = wb.create_sheet("QA")
    qa.append(["Sheet", "Row", "Issue"])
    n = 0
    seen = set()  # document-wide: a repeated line on another page is exactly what a client wants flagged
    for pno, page in enumerate(doc, 1):
        for t in page.find_tables().tables:
            rows = [[clean(c) for c in r] for r in t.extract()]
            rows = [r for r in rows if any(c not in (None, "") for c in r)]
            if len(rows) < 2:
                continue
            n += 1
            ws = wb.create_sheet(f"Table{n}_p{pno}")
            for i, r in enumerate(rows, 1):
                ws.append(r)
                key = tuple(r)
                if i > 1 and key in seen:
                    qa.append([ws.title, i, "duplicate row"])
                seen.add(key)
                for j, c in enumerate(r, 1):
                    if i > 1 and c in (None, ""):
                        qa.append([ws.title, i, f"empty cell in column {get_column_letter(j)}"])
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="DDEBF7")
            for col in ws.columns:
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(45, max(10, max(len(str(c.value or "")) for c in col) + 2))
            ws.freeze_panes = "A2"
            summary.append([ws.title, pno, len(rows) - 1, len(rows[0])])
    for c in summary[1]:
        c.font = Font(bold=True)
    wb.save(xlsx_path)
    return n


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    print(f"{convert(src, dst)} table(s) -> {dst}")
