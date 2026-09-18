"""Builds sales_dashboard.xlsx: a Data sheet (360 synthetic orders), a Dashboard sheet driven entirely by live formulas
(SUMIFS / COUNTIFS / AVERAGEIFS), a monthly trend chart and a category bar chart. Change the data, the dashboard updates."""
import pathlib
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(11)
wb = Workbook()
data = wb.active
data.title = "Data"
data.append(["Order ID", "Date", "Month", "Region", "Category", "Units", "Unit price", "Revenue", "Channel"])
regions, cats, channels = ["North", "South", "East", "West"], ["Hardware", "Software", "Services", "Training"], ["Web", "Partner", "Direct"]
for i in range(360):
    m = random.randint(1, 9)
    d = f"2026-{m:02d}-{random.randint(1, 28):02d}"
    units, price = random.randint(1, 25), random.choice([49, 99, 199, 349, 799])
    r = i + 2
    data.append([f"SO-{10000 + i}", d, None, random.choice(regions), random.choice(cats), units, price, None, random.choice(channels)])
    data[f"C{r}"] = f'=TEXT(DATEVALUE(B{r}),"yyyy-mm")'
    data[f"H{r}"] = f"=F{r}*G{r}"
for c in data[1]:
    c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
for i, w in enumerate([11, 12, 10, 9, 11, 7, 11, 11, 9], 1):
    data.column_dimensions[get_column_letter(i)].width = w
data.freeze_panes = "A2"
data.auto_filter.ref = f"A1:I{data.max_row}"

db = wb.create_sheet("Dashboard", 0)
db["A1"] = "Sales dashboard — live formulas over the Data sheet"
db["A1"].font = Font(bold=True, size=14)
kpis = [("Total revenue", "=SUM(Data!H:H)"), ("Orders", "=COUNTA(Data!A:A)-1"), ("Avg order value", "=AVERAGE(Data!H2:H1000)"),
        ("Units sold", "=SUM(Data!F:F)")]
for i, (k, f) in enumerate(kpis):
    col = get_column_letter(1 + i * 2)
    db[f"{col}3"], db[f"{col}4"] = k, f
    db[f"{col}3"].font, db[f"{col}4"].font = Font(color="808080"), Font(bold=True, size=13)
    db[f"{col}4"].number_format = '#,##0'

db["A7"], db["B7"], db["C7"] = "Month", "Revenue", "Orders"
months = [f"2026-{m:02d}" for m in range(1, 10)]
for i, m in enumerate(months, 8):
    db[f"A{i}"] = m
    db[f"B{i}"] = f'=SUMIFS(Data!H:H,Data!C:C,A{i})'
    db[f"C{i}"] = f'=COUNTIFS(Data!C:C,A{i})'
db["E7"], db["F7"], db["G7"] = "Category", "Revenue", "Avg units"
for i, c in enumerate(cats, 8):
    db[f"E{i}"] = c
    db[f"F{i}"] = f'=SUMIFS(Data!H:H,Data!E:E,E{i})'
    db[f"G{i}"] = f'=AVERAGEIFS(Data!F:F,Data!E:E,E{i})'
db["I7"], db["J7"] = "Region", "Revenue"
for i, rg in enumerate(regions, 8):
    db[f"I{i}"], db[f"J{i}"] = rg, f'=SUMIFS(Data!H:H,Data!D:D,I{i})'
for cell in ("A7", "B7", "C7", "E7", "F7", "G7", "I7", "J7"):
    db[cell].font, db[cell].fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
for rng in ("B8:B16", "F8:F11", "J8:J11"):
    for row in db[rng]:
        for c in row:
            c.number_format = '#,##0'

line = LineChart()
line.title, line.height, line.width = "Monthly revenue", 7, 16
line.add_data(Reference(db, min_col=2, min_row=7, max_row=16), titles_from_data=True)
line.set_categories(Reference(db, min_col=1, min_row=8, max_row=16))
db.add_chart(line, "A19")
bar = BarChart()
bar.title, bar.height, bar.width = "Revenue by category", 7, 12
bar.add_data(Reference(db, min_col=6, min_row=7, max_row=11), titles_from_data=True)
bar.set_categories(Reference(db, min_col=5, min_row=8, max_row=11))
db.add_chart(bar, "I19")
for col, w in zip("ABCDEFGHIJ", [12, 12, 9, 3, 12, 12, 10, 3, 10, 12]):
    db.column_dimensions[col].width = w
out = pathlib.Path(__file__).parent / "sales_dashboard.xlsx"
wb.save(out)
print(out)
